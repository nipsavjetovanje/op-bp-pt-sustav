"""Renderira strukturirane PT podatke (iz pt_generator.generate_pt) u profesionalni PT.docx."""


def _fmt_num(n):
    f = float(n or 0)
    return str(int(f)) if f == int(f) else f"{f:.1f}"

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x38, 0x64)
GRAY = RGBColor(0x59, 0x59, 0x59)
LIGHT_BG = "EEF2F7"


def _shade_cell(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def _set_cell_text(cell, text, bold=False, color=None, size=9, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def _header_row(table, headers, widths):
    row = table.rows[0]
    for i, text in enumerate(headers):
        cell = row.cells[i]
        _set_cell_text(cell, text, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=9)
        _shade_cell(cell, "1F3864")


def render_pt_docx(
    pt_data: dict,
    output_path: str,
    company_name: str = "NIKOLIĆ I PARTNERI",
    trajanje_mjeseci: float | None = None,
    sati_mjesecno: float = 169.0,
):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    # Letterhead (header)
    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    run = hp.add_run(company_name)
    run.bold = True
    run.font.color.rgb = NAVY
    run.font.size = Pt(10)

    # Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("Nikolić i partneri j.d.o.o. — Računovodstvo i EU fondovi — Zagreb")
    frun.font.size = Pt(7)
    frun.font.color.rgb = GRAY

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    trun = title.add_run("PROJEKTNI TIM (PT)")
    trun.bold = True
    trun.font.size = Pt(18)
    trun.font.color.rgb = NAVY

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    srun = subtitle.add_run("Generirano automatski iz CV-a, platne liste i metodologije izračuna satnice")
    srun.italic = True
    srun.font.size = Pt(9)
    srun.font.color.rgb = GRAY

    if pt_data.get("role_vocabulary_used"):
        note = doc.add_paragraph()
        nrun = note.add_run("Napomena: uloge su usklađene s propisanim vokabularom iz UP/BL natječajne dokumentacije.")
        nrun.italic = True
        nrun.font.size = Pt(8.5)
        nrun.font.color.rgb = GRAY

    # 1. Metodologija
    h1 = doc.add_paragraph()
    r = h1.add_run("1. Metodologija izračuna satnice")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = NAVY

    quote_table = doc.add_table(rows=1, cols=1)
    quote_table.style = "Table Grid"
    cell = quote_table.rows[0].cells[0]
    _set_cell_text(cell, pt_data.get("methodology_quote", "(nije generirano)"), italic=True, size=9.5)

    note1 = doc.add_paragraph()
    n1run = note1.add_run(
        "Ova sekcija je doslovan citat natječajnog priloga koji propisuje izračun satnice — satnice "
        "u nastavku moraju biti sljedive na ovaj tekst."
    )
    n1run.italic = True
    n1run.font.size = Pt(8.5)
    n1run.font.color.rgb = GRAY

    members = pt_data.get("members", [])

    # 2. Pregled tima
    h2 = doc.add_paragraph()
    r = h2.add_run("2. Pregled članova projektnog tima")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = NAVY

    t2 = doc.add_table(rows=1 + max(len(members), 1), cols=3)
    t2.style = "Table Grid"
    _header_row(t2, ["Ime i prezime", "Uloga u projektu", "Ključne kompetencije (iz CV-a)"], None)
    for i, m in enumerate(members, start=1):
        row = t2.rows[i]
        _set_cell_text(row.cells[0], m.get("ime", ""))
        _set_cell_text(row.cells[1], m.get("uloga", ""))
        _set_cell_text(row.cells[2], m.get("kompetencije", ""))
    if not members:
        _set_cell_text(t2.rows[1].cells[0], "(nema generiranih članova)", italic=True, color=GRAY)

    # 3. Opis rada po aktivnostima
    h3 = doc.add_paragraph()
    r = h3.add_run("3. Opis rada po aktivnostima")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = NAVY

    rows_data = []
    for m in members:
        for a in m.get("aktivnosti", []):
            rows_data.append((m.get("ime", ""), a))

    t3 = doc.add_table(rows=1 + max(len(rows_data), 1), cols=6)
    t3.style = "Table Grid"
    _header_row(t3, ["Član tima", "Aktivnost", "Opis rada", "Sat rada", "Satnica (EUR/h)", "Trošak (EUR)"], None)
    total_cost = 0.0
    for i, (ime, a) in enumerate(rows_data, start=1):
        row = t3.rows[i]
        sat = float(a.get("sat_rada", 0) or 0)
        satnica = round(float(a.get("satnica", 0) or 0), 2)
        trosak = round(sat * satnica, 2)
        total_cost += trosak
        _set_cell_text(row.cells[0], ime)
        _set_cell_text(row.cells[1], a.get("aktivnost", ""))
        _set_cell_text(row.cells[2], a.get("opis_rada", ""))
        _set_cell_text(row.cells[3], _fmt_num(sat))
        _set_cell_text(row.cells[4], f"{satnica:.2f}")
        _set_cell_text(row.cells[5], f"{trosak:.2f}")
    if not rows_data:
        _set_cell_text(t3.rows[1].cells[0], "(nema podataka)", italic=True, color=GRAY)

    note3 = doc.add_paragraph()
    n3run = note3.add_run(
        "Sat rada procijenjen je prema opsegu aktivnosti (ne prema raspoloživom budžetu). "
        "Satnica je zaokružena na 2 decimale."
    )
    n3run.italic = True
    n3run.font.size = Pt(8.5)
    n3run.font.color.rgb = GRAY

    # 4. Rekapitulacija
    h4 = doc.add_paragraph()
    r = h4.add_run("4. Rekapitulacija troška tima")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = NAVY

    per_member = {}
    for ime, a in rows_data:
        sat = float(a.get("sat_rada", 0) or 0)
        satnica = round(float(a.get("satnica", 0) or 0), 2)
        d = per_member.setdefault(ime, {"sati": 0.0, "trosak": 0.0})
        d["sati"] += sat
        d["trosak"] += round(sat * satnica, 2)

    raspoloziv_sati = (trajanje_mjeseci or 0) * sati_mjesecno

    t4 = doc.add_table(rows=1 + len(per_member) + 1, cols=4)
    t4.style = "Table Grid"
    _header_row(t4, ["Član tima", "Ukupno sati", "Ukupan trošak (EUR)", "% rada na projektu (FTE)"], None)
    i = 1
    sum_sati = 0.0
    sum_trosak = 0.0
    for ime, d in per_member.items():
        row = t4.rows[i]
        _set_cell_text(row.cells[0], ime)
        _set_cell_text(row.cells[1], _fmt_num(d["sati"]))
        _set_cell_text(row.cells[2], f"{d['trosak']:.2f}")
        postotak = (d["sati"] / raspoloziv_sati * 100) if raspoloziv_sati > 0 else None
        _set_cell_text(row.cells[3], f"{postotak:.1f}%" if postotak is not None else "—")
        sum_sati += d["sati"]
        sum_trosak += d["trosak"]
        i += 1
    total_row = t4.rows[i]
    _set_cell_text(total_row.cells[0], "UKUPNO", bold=True)
    _set_cell_text(total_row.cells[1], _fmt_num(sum_sati), bold=True)
    _set_cell_text(total_row.cells[2], f"{sum_trosak:.2f}", bold=True)
    _set_cell_text(total_row.cells[3], "", bold=True)

    note4 = doc.add_paragraph()
    note4_text = "Ovaj dokument je izvor podataka za redove troška osoblja u obrascu BP i za sažetak tima u obrascu OP."
    if raspoloziv_sati > 0:
        note4_text += (
            f" % rada na projektu (FTE) računa se kao ukupno dodijeljeni sati / ({sati_mjesecno:.0f} h × "
            f"{trajanje_mjeseci:.0f} mjeseci trajanja projekta)."
        )
    else:
        note4_text += " Trajanje projekta nije uneseno — % rada na projektu nije izračunat."
    n4run = note4.add_run(note4_text)
    n4run.italic = True
    n4run.font.size = Pt(8.5)
    n4run.font.color.rgb = GRAY

    doc.save(output_path)
    return {"ukupno_sati": sum_sati, "ukupan_trosak": round(sum_trosak, 2)}
