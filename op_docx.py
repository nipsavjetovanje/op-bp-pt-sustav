"""Renderira strukturirane OP podatke (iz op_generator.generate_op) u OP.docx,
u formatu kompatibilnom s primjerom OP obrasca (Naziv, Opis poslovanja/tržišta,
Opis budućeg projekta, Ciljano tržište, Faze+Aktivnosti, Projektni tim, Oprema)."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x38, 0x64)
GRAY = RGBColor(0x59, 0x59, 0x59)


def _shade_cell(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color_hex)
    tcPr.append(shd)


def _field_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    return p


def _field_box(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text or "")
    run.font.size = Pt(10.5)
    return table


def _activity_list_box(doc, activities):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.rows[0].cells[0]
    cell.text = ""
    for i, a in enumerate(activities):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        run = p.add_run(f"{i + 1}. {a}")
        run.font.size = Pt(10.5)
    return table


def render_op_docx(op_data: dict, output_path: str, company_name: str = "NIKOLIĆ I PARTNERI"):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    header = section.header
    hp = header.paragraphs[0]
    run = hp.add_run(company_name)
    run.bold = True
    run.font.color.rgb = NAVY
    run.font.size = Pt(10)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("Nikolić i partneri j.d.o.o. — Računovodstvo i EU fondovi — Zagreb")
    frun.font.size = Pt(7)
    frun.font.color.rgb = GRAY

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    trun = title.add_run("2. Opis projekta (OP)")
    trun.bold = True
    trun.font.size = Pt(18)
    trun.font.color.rgb = NAVY

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    srun = subtitle.add_run("Generirano automatski iz upitnika, popisa aktivnosti i PT obrasca")
    srun.italic = True
    srun.font.size = Pt(9)
    srun.font.color.rgb = GRAY

    _field_heading(doc, "Naziv projekta:")
    _field_box(doc, op_data.get("naziv_projekta", ""))

    _field_heading(doc, "Opis poslovanja / tržišta:")
    _field_box(doc, op_data.get("opis_poslovanja", ""))

    _field_heading(doc, "Opis budućeg projekta (djelatnost, proširenje):")
    _field_box(doc, op_data.get("opis_buduceg_projekta", ""))

    _field_heading(doc, "Ciljano tržište:")
    _field_box(doc, op_data.get("ciljano_trziste", ""))

    _field_heading(doc, "Faze projekta i Aktivnosti projekta:")
    _field_box(doc, op_data.get("faze_i_aktivnosti_uvod", ""))

    p = doc.add_paragraph()
    r = p.add_run("Aktivnosti projekta:")
    r.bold = True
    r.font.size = Pt(11)
    _activity_list_box(doc, op_data.get("aktivnosti_lista", []))
    note = doc.add_paragraph()
    nrun = note.add_run("Aktivnosti koje će se koristiti u sklopu ovog projekta i koje su sastavni dio dokumenta UP.")
    nrun.italic = True
    nrun.font.size = Pt(8.5)
    nrun.font.color.rgb = GRAY

    _field_heading(doc, "Projektni tim:")
    _field_box(doc, op_data.get("projektni_tim_sazetak", ""))
    note2 = doc.add_paragraph()
    n2run = note2.add_run("Detaljan opis rada po aktivnosti i troškovi nalaze se u obrascu PT.")
    n2run.italic = True
    n2run.font.size = Pt(8.5)
    n2run.font.color.rgb = GRAY

    _field_heading(doc, "Oprema prijavitelja (postojeća):")
    _field_box(doc, op_data.get("oprema", ""))

    doc.save(output_path)
