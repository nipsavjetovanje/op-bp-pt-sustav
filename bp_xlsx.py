"""Renderira strukturirane BP podatke (iz bp_generator.generate_bp) u BP.xlsx."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY = "1F3864"
LIGHT_RED = "FFE5E5"

HEADERS = [
    "Kategorija troška", "Opis troška/Naziv", "Količina", "Jedinična cijena",
    "Ukupno", "Postotak financiranja", "Bespovratna sredstva", "Vlastito učešće",
]


def render_bp_xlsx(bp_data: dict, output_path: str, warnings: list[str] | None = None):
    wb = Workbook()
    ws = wb.active
    ws.title = "BP"

    header_fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    stavke = bp_data.get("stavke", [])
    row_idx = 2
    for s in stavke:
        ws.cell(row=row_idx, column=1, value=s["kategorija"]).border = border
        ws.cell(row=row_idx, column=2, value=s["opis_naziv"]).border = border
        ws.cell(row=row_idx, column=3, value=s["kolicina"]).border = border
        ws.cell(row=row_idx, column=4, value=s["jedinicna_cijena"]).border = border
        ws.cell(row=row_idx, column=5, value=s["ukupno"]).border = border
        ws.cell(row=row_idx, column=6, value=f"{s['postotak_financiranja']:.0f}%").border = border
        ws.cell(row=row_idx, column=7, value=s["bespovratna_sredstva"]).border = border
        ws.cell(row=row_idx, column=8, value=s["vlastito_ucesce"]).border = border
        row_idx += 1

    # Red UKUPNO
    total_row = row_idx
    ws.cell(row=total_row, column=2, value="UKUPNO").font = Font(bold=True)
    for col, key in [(5, "ukupno"), (7, "bespovratna_sredstva"), (8, "vlastito_ucesce")]:
        total = sum(s[key] for s in stavke)
        c = ws.cell(row=total_row, column=col, value=round(total, 2))
        c.font = Font(bold=True)
        c.border = border
    row_idx += 2

    # Rekapitulacija po kategoriji
    ws.cell(row=row_idx, column=1, value="Rekapitulacija po kategoriji").font = Font(bold=True, size=11, color=NAVY)
    row_idx += 1
    by_cat = {}
    for s in stavke:
        by_cat.setdefault(s["kategorija"], 0.0)
        by_cat[s["kategorija"]] += s["ukupno"]
    total_all = sum(by_cat.values()) or 1
    limits = {k["naziv"]: k.get("max_udio_posto") for k in bp_data.get("kategorije", [])}
    for kat, iznos in by_cat.items():
        udio = iznos / total_all * 100
        limit = limits.get(kat)
        flag = limit is not None and udio > limit + 0.01
        ws.cell(row=row_idx, column=1, value=kat)
        ws.cell(row=row_idx, column=2, value=f"{iznos:.2f} EUR ({udio:.1f}%)")
        limit_text = f"limit: {limit:.1f}%" if limit is not None else "limit nije prepoznat u UP-u"
        c = ws.cell(row=row_idx, column=3, value=limit_text)
        if flag:
            for col in (1, 2, 3):
                ws.cell(row=row_idx, column=col).fill = PatternFill(
                    start_color=LIGHT_RED, end_color=LIGHT_RED, fill_type="solid"
                )
        row_idx += 1

    row_idx += 1
    max_iznos = bp_data.get("max_iznos_potpore")
    total_bespovratna = sum(s["bespovratna_sredstva"] for s in stavke)
    max_text = f"{max_iznos:.2f} EUR" if max_iznos is not None else "nije prepoznat u UP-u"
    label = f"Max. iznos potpore (UP): {max_text}   |   Traženo: {total_bespovratna:.2f} EUR"
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=8)
    cell = ws.cell(row=row_idx, column=1, value=label)
    cell.font = Font(bold=True)
    if max_iznos is not None and total_bespovratna > max_iznos + 0.01:
        cell.fill = PatternFill(start_color=LIGHT_RED, end_color=LIGHT_RED, fill_type="solid")
    row_idx += 2

    if warnings:
        ws.cell(row=row_idx, column=1, value="Upozorenja post-validacije").font = Font(bold=True, size=11, color=NAVY)
        row_idx += 1
        for w in warnings:
            ws.cell(row=row_idx, column=1, value=w)
            row_idx += 1

    widths = [30, 40, 10, 14, 12, 16, 16, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.freeze_panes = "A2"

    wb.save(output_path)
